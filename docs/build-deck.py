"""Build the Earthen client walkthrough deck. Run: python3 docs/build-deck.py"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

CREAM = RGBColor(0xF7, 0xF4, 0xEE)
PAPER = RGBColor(0xFB, 0xF9, 0xF5)
INK = RGBColor(0x1C, 0x1A, 0x17)
MUTED = RGBColor(0x6B, 0x65, 0x60)
LINE = RGBColor(0xE4, 0xDD, 0xD2)
OLIVE = RGBColor(0x6E, 0x76, 0x58)
OLIVE_DEEP = RGBColor(0x4F, 0x56, 0x38)
TERRA = RGBColor(0xA3, 0x4B, 0x36)
NIGHT = RGBColor(0x2A, 0x29, 0x26)
SAND = RGBColor(0xE4, 0xD7, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W = Inches(13.333)
H = Inches(7.5)


def rgb_hex(color):
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_run(run, text, font, size, color, bold=False, italic=False, spacing=None):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))


def add_text(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
    frame._txBody.bodyPr.set("anchor", anchor_map[anchor])
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = line.get("align", PP_ALIGN.LEFT)
        paragraph.space_after = Pt(line.get("after", 0))
        paragraph.space_before = Pt(line.get("before", 0))
        if "level" in line:
            paragraph.level = line["level"]
        run = paragraph.add_run()
        set_run(
            run,
            line["text"],
            line.get("font", "Calibri"),
            line.get("size", 16),
            line.get("color", INK),
            line.get("bold", False),
            line.get("italic", False),
            line.get("spacing"),
        )
    return shape


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CREAM
    return slide


def chrome(slide, kicker, page, pages):
    rect(slide, 0, 0, 13.333, 0.08, OLIVE)
    add_text(
        slide,
        0.62,
        0.28,
        8,
        0.32,
        [{"text": kicker.upper(), "size": 12, "color": OLIVE, "bold": True, "spacing": 2.2}],
    )
    rect(slide, 0.62, 7.12, 12.1, 0.01, LINE)
    add_text(slide, 0.62, 7.16, 4, 0.26, [{"text": "EARTHEN", "size": 10, "color": MUTED, "spacing": 1.8}])
    add_text(
        slide,
        10.4,
        7.16,
        2.3,
        0.26,
        [{"text": f"{page}  /  {pages}", "size": 10, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )


def title_block(slide, title, subtitle=None):
    add_text(slide, 0.62, 0.58, 12, 0.7, [{"text": title, "font": "Georgia", "size": 32, "color": INK}])
    if subtitle:
        add_text(slide, 0.62, 1.28, 11.5, 0.4, [{"text": subtitle, "size": 15, "color": MUTED}])


def bullet_block(slide, x, y, w, h, items, size=15, color=INK):
    lines = []
    for item in items:
        lines.append({"text": item, "size": size, "color": color, "after": 9})
    return add_text(slide, x, y, w, h, lines)


PAGES = 16
HERE = Path(__file__).resolve().parent
SHOTS = HERE / "slides"


def picture(slide, filename, x, y, w, h):
    pad = 0.03
    rect(slide, x - pad, y - pad, w + pad * 2, h + pad * 2, WHITE, LINE)
    return slide.shapes.add_picture(str(SHOTS / filename), Inches(x), Inches(y), Inches(w), Inches(h))


def live(prs, where, title, filename, caption):
    """A full screenshot from the site running on this machine."""
    slide = blank(prs)
    rect(slide, 0, 0, 13.333, 0.08, OLIVE)
    add_text(
        slide,
        0.48,
        0.2,
        1.1,
        0.26,
        [{"text": "LIVE", "size": 11, "color": OLIVE, "bold": True, "spacing": 2.4}],
    )
    add_text(slide, 1.45, 0.18, 8, 0.28, [{"text": where, "size": 13, "color": MUTED}])
    pic_h = 5.62
    pic_w = pic_h * 1440 / 900
    picture(slide, filename, (13.333 - pic_w) / 2, 0.56, pic_w, pic_h)
    add_text(slide, 0.48, 6.36, 9.6, 0.32, [{"text": title, "font": "Georgia", "size": 16}])
    add_text(slide, 0.48, 6.7, 10.4, 0.36, [{"text": caption, "size": 13, "color": MUTED}])
    add_text(
        slide,
        9.5,
        7.16,
        3.35,
        0.26,
        [{"text": "localhost:3000", "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Earthen — how the website works"
    prs.core_properties.subject = "Client walkthrough of the shop and the studio"
    prs.core_properties.author = "Earthen"

    title(prs)
    two_sides(prs)
    path(prs)
    home(prs)
    live(prs, "The shop", "Arriving on the shop", "home.png", "The homepage as it is running now, with the announcement the studio has set.")
    shop(prs)
    live(prs, "The shop", "The clothing grid", "shop.png", "All clothing, newest first, with size, colour, gender, and stock.")
    product(prs)
    live(prs, "The shop", "Linen Overshirt", "product.png", "Colour, size, a delivery check, add to bag, and the wishlist.")
    checkout(prs)
    live(prs, "The shop", "The bag", "cart.png", "Three pieces, quantities, and a place for a code.")
    live(prs, "The shop", "Checkout", "checkout.png", "Aisha is signed in. On this machine the page still says no payment is taken.")
    after(prs)
    live(prs, "The shop", "The order, once it is placed", "order.png", "EA-10424 is confirmed. The invoice and cancel are both here.")
    account(prs)
    live(prs, "The shop", "A signed-in account", "account.png", "Orders, the wishlist, a saved address, and what was viewed last.")
    roles(prs)
    dashboard(prs)
    live(prs, "The studio", "The morning view", "admin-dashboard.png", "Revenue today, open shipments, low stock, and the latest orders.")
    catalog(prs)
    live(prs, "The studio", "The catalog", "admin-products.png", "Search and filters, a spreadsheet import, and add product.")
    orders(prs)
    live(prs, "The studio", "The order list", "admin-orders.png", "Search, status filters, and export.")
    desk(prs)
    live(prs, "The studio", "Customers", "admin-customers.png", "Orders, amount spent, and who is active. Guest checkouts appear by email.")
    live(prs, "The studio", "Coupons", "admin-coupons.png", "EARTHEN200, FREESHIP, and WELCOME10, each with a use limit.")
    reports(prs)
    live(prs, "The studio", "Sales and GST", "admin-reports.png", "This period’s sales, how people paid, and the October GST lines.")
    settings(prs)
    live(prs, "The studio", "What the owner can change", "admin-settings.png", "The announcement, staff, and the store-wide sale. Nothing reaches the shop until Save changes.")
    letters(prs)
    close(prs)
    notes(prs)

    out = HERE / "Earthen-how-the-website-works.pptx"
    prs.save(out)
    return out


def title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NIGHT
    rect(slide, 0, 0, 0.18, 7.5, OLIVE)
    add_text(
        slide,
        0.85,
        1.85,
        10,
        0.35,
        [{"text": "EARTHEN", "size": 14, "color": SAND, "spacing": 4.5}],
    )
    add_text(
        slide,
        0.85,
        2.3,
        11,
        1.6,
        [{"text": "How the website works", "font": "Georgia", "size": 48, "color": CREAM}],
    )
    add_text(
        slide,
        0.85,
        4.15,
        9,
        0.9,
        [
            {"text": "A walkthrough of the shop and the studio,", "size": 18, "color": SAND, "after": 4},
            {"text": "from the first visit to sign-out.", "size": 18, "color": SAND},
        ],
    )
    add_text(slide, 0.85, 6.55, 8, 0.3, [{"text": "Quiet clothes for every day.", "size": 14, "color": OLIVE, "italic": True, "font": "Georgia"}])


def two_sides(prs):
    slide = blank(prs)
    chrome(slide, "The website", 1, PAGES)
    title_block(slide, "Two doors, two logins", "A shopper and a member of staff never share a sign-in.")
    rect(slide, 0.62, 2.05, 5.85, 4.45, PAPER, LINE)
    rect(slide, 0.62, 2.05, 0.1, 4.45, TERRA)
    add_text(slide, 1.0, 2.3, 5, 0.3, [{"text": "THE SHOP", "size": 12, "color": TERRA, "bold": True, "spacing": 1.8}])
    add_text(slide, 1.0, 2.7, 5, 0.6, [{"text": "For customers", "font": "Georgia", "size": 28}])
    bullet_block(
        slide,
        1.0,
        3.55,
        5.1,
        2.6,
        [
            "Home, search, a piece, the bag, checkout, account.",
            "Anyone can browse. Checkout needs an email, not an account.",
            "An account keeps orders, addresses, the bag, and the wishlist.",
            "Prices are in rupees and include GST.",
        ],
        15,
    )
    rect(slide, 6.85, 2.05, 5.85, 4.45, OLIVE_DEEP)
    add_text(slide, 7.2, 2.3, 5, 0.3, [{"text": "THE STUDIO", "size": 12, "color": SAND, "bold": True, "spacing": 1.8}])
    add_text(slide, 7.2, 2.7, 5.1, 0.6, [{"text": "For staff", "font": "Georgia", "size": 28, "color": CREAM}])
    bullet_block(
        slide,
        7.2,
        3.55,
        5.1,
        2.6,
        [
            "A separate sign-in at the studio door.",
            "Owner, catalog manager, or order and support.",
            "Each role sees only its own pages.",
            "Orders, stock, coupons, tax, and the look of the shop.",
        ],
        15,
        CREAM,
    )


def path(prs):
    slide = blank(prs)
    chrome(slide, "The visit", 2, PAGES)
    title_block(slide, "What we will walk through", "Seven steps for the customer. Staff pick the order up at step four.")
    steps = [
        ("01", "Arrive", "Home, announcement, and the promises."),
        ("02", "Find a piece", "Search, filters, and the grid."),
        ("03", "Choose", "Colour, size, bag, or wishlist."),
        ("04", "Pay", "Address, delivery, coupon, and payment."),
        ("05", "Follow", "Confirmed through to delivered."),
        ("06", "Aftercare", "Cancel, return, invoice, or a review."),
        ("07", "Account", "Profile, orders, wishlist, sign-out."),
    ]
    for index, (num, name, text) in enumerate(steps):
        if index < 4:
            x = 0.62 + index * 3.15
            y = 2.0
        else:
            x = 2.2 + (index - 4) * 3.15
            y = 4.25
        rect(slide, x, y, 2.95, 2.05, PAPER, LINE)
        add_text(slide, x + 0.2, y + 0.18, 2.5, 0.3, [{"text": num, "size": 12, "color": TERRA, "bold": True, "spacing": 1.2}])
        add_text(slide, x + 0.2, y + 0.52, 2.5, 0.4, [{"text": name, "font": "Georgia", "size": 20}])
        add_text(slide, x + 0.2, y + 1.05, 2.55, 0.75, [{"text": text, "size": 13, "color": MUTED}])


def home(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 3, PAGES)
    title_block(slide, "Arriving on the shop")
    bullet_block(
        slide,
        0.62,
        1.7,
        7.3,
        4.9,
        [
            "A season name, a headline, and two ways in: Shop the collection, and Our story.",
            "Hero photographs. A live store-wide sale shows its name and the percent off.",
            "Women, Men, and Unisex, then categories, including New in.",
            "A season banner, new arrivals, and up to four recently viewed pieces.",
            "An olive announcement bar. Staff write it, or the site fills in shipping and returns. It can be hidden.",
            "The header carries search, account, and the bag count. On a phone it also opens Our story and Help.",
            "The footer holds policies, contact, the studio address, and GSTIN.",
        ],
        16,
    )
    promises = [
        ("Free shipping", "Over the studio threshold"),
        ("Returns", "A set window, pickup arranged"),
        ("Cash on delivery", "When the studio allows it"),
        ("Secure payment", "UPI, cards, net banking"),
    ]
    for index, (name, text) in enumerate(promises):
        y = 1.7 + index * 1.2
        rect(slide, 8.25, y, 4.45, 1.05, PAPER, LINE)
        add_text(slide, 8.48, y + 0.16, 4.0, 0.32, [{"text": name, "size": 16, "bold": True}])
        add_text(slide, 8.48, y + 0.5, 4.0, 0.35, [{"text": text, "size": 13, "color": MUTED}])


def shop(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 4, PAGES)
    title_block(slide, "Finding a piece")
    rect(slide, 0.62, 1.85, 6.0, 4.85, PAPER, LINE)
    add_text(slide, 0.88, 2.05, 5.4, 0.35, [{"text": "Search", "font": "Georgia", "size": 22}])
    bullet_block(
        slide,
        0.88,
        2.6,
        5.45,
        3.8,
        [
            "Opens over the page as they type.",
            "Matches show the price. A sale price is crossed out.",
            "Popular searches: linen, overshirt, trousers, cardigan, olive.",
            "Looks through name, fabric, origin, description, category, and colour.",
            "A one-letter miss still matches.",
        ],
        15,
    )
    rect(slide, 6.85, 1.85, 5.85, 4.85, PAPER, LINE)
    add_text(slide, 7.1, 2.05, 5.4, 0.35, [{"text": "The grid", "font": "Georgia", "size": 22}])
    bullet_block(
        slide,
        7.1,
        2.6,
        5.3,
        3.8,
        [
            "Twenty pieces a page, from a category, New in, or an audience.",
            "Sort by newest, popularity, or price.",
            "Filter size, colour, gender, price, and stock.",
            "Women and Men also keep unisex pieces.",
            "Save or add to the bag from the card, without opening the piece.",
        ],
        15,
    )


def product(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 5, PAGES)
    title_block(slide, "On a piece", "Drafts and archived pieces stay off the shop. Only active pieces can be bought.")
    left = [
        "Photographs, price, fabric, stars, and a New in mark.",
        "Who it was photographed on, and how it fits.",
        "Colour and size. A saved usual size is pre-selected.",
        "Gone sizes are struck through. Three or fewer left is said aloud on the page.",
        "A size guide in centimetres, upper or lower body.",
    ]
    right = [
        "Details, origin, fabric, care, shipping, and returns.",
        "Add to bag. A panel offers View bag or Checkout.",
        "Wishlist, and remove it the same way.",
        "Out of stock: leave an email and be written to when it returns.",
        "A review after delivery: stars, fit, and an optional height.",
        "A PIN check, and four related pieces underneath.",
    ]
    bullet_block(slide, 0.62, 1.9, 6.0, 4.8, left, 15)
    bullet_block(slide, 6.9, 1.9, 5.8, 4.8, right, 15)


def checkout(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 6, PAGES)
    title_block(slide, "Bag, then checkout", "No account is required. The receipt goes to the email on the order.")
    rect(slide, 0.62, 1.9, 6.0, 4.8, PAPER, LINE)
    add_text(slide, 0.88, 2.1, 5.4, 0.35, [{"text": "The bag", "font": "Georgia", "size": 22}])
    bullet_block(
        slide,
        0.88,
        2.65,
        5.45,
        3.7,
        [
            "Colour, size, quantity, and remove.",
            "Type a code, or open Offers for this order.",
            "Each offer shows the saving, the minimum, and why it may not apply.",
            "A bar shows how close they are to free standard delivery.",
            "Subtotal, shipping, discount, and a GST-inclusive total.",
        ],
        15,
    )
    rect(slide, 6.85, 1.9, 5.85, 4.8, PAPER, LINE)
    add_text(slide, 7.1, 2.1, 5.3, 0.35, [{"text": "Checkout", "font": "Georgia", "size": 22}])
    bullet_block(
        slide,
        7.1,
        2.65,
        5.3,
        3.7,
        [
            "Address, a serviceable PIN, and saved addresses when signed in.",
            "Standard or express. A free-shipping code waives either.",
            "Optional gift wrap, a card note, and prices left off the packing slip.",
            "UPI, card, or net banking through Razorpay.",
            "Cash on delivery when it is on, and under the studio limit.",
            "Stock drops as the order is placed. The bag empties.",
        ],
        15,
    )


def after(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 7, PAGES)
    title_block(slide, "After the order is placed")
    stages = ["Confirmed", "Packed", "Shipped", "Out for delivery", "Delivered"]
    for index, name in enumerate(stages):
        x = 0.62 + index * 2.5
        rect(slide, x, 1.85, 2.3, 0.85, OLIVE if index < 4 else TERRA)
        add_text(
            slide,
            x + 0.08,
            2.05,
            2.14,
            0.5,
            [{"text": name, "size": 13, "color": CREAM, "bold": True, "align": PP_ALIGN.CENTER}],
            anchor=MSO_ANCHOR.MIDDLE,
        )
    cards = [
        ("Follow it", "Status, courier tracking once staff add it, payment method, and the city it is going to."),
        ("Change it", "Cancel while it is confirmed or packed. Stock returns. A paid Razorpay order is sent for refund."),
        ("Bring it back", "After delivery, inside the studio’s return window: refund or exchange, with pickup from the same address."),
        ("Keep the paper", "A GST invoice to print or save. Review the piece once it has arrived."),
    ]
    for index, (name, text) in enumerate(cards):
        col = index % 2
        row = index // 2
        x = 0.62 + col * 6.35
        y = 3.05 + row * 1.9
        rect(slide, x, y, 6.1, 1.72, PAPER, LINE)
        add_text(slide, x + 0.25, y + 0.18, 5.6, 0.35, [{"text": name, "font": "Georgia", "size": 18}])
        add_text(slide, x + 0.25, y + 0.65, 5.6, 0.85, [{"text": text, "size": 14, "color": MUTED}])


def account(prs):
    slide = blank(prs)
    chrome(slide, "The shop", 8, PAGES)
    title_block(slide, "The account, from sign-in to sign-out")
    items = [
        ("Create", "Name, email, a 10-digit mobile, and a password of at least 8 characters."),
        ("Sign in", "Email and password. A blocked account is refused. Forgot password sends a link that lasts 30 minutes."),
        ("Overview", "Orders, wishlist, addresses, the latest order, and pieces they asked to hear about."),
        ("Profile", "Name and mobile. A usual size. An address book, with one default."),
        ("Orders", "Number, date, status, and total. Guest orders appear when the email matches."),
        ("Sign out", "Ends the session. The shop stays open. The wishlist on a signed-out browser stays on that browser."),
    ]
    for index, (name, text) in enumerate(items):
        col = index % 3
        row = index // 3
        x = 0.62 + col * 4.15
        y = 1.9 + row * 2.4
        rect(slide, x, y, 3.95, 2.2, PAPER, LINE)
        add_text(slide, x + 0.22, y + 0.2, 3.5, 0.4, [{"text": name, "font": "Georgia", "size": 20}])
        add_text(slide, x + 0.22, y + 0.75, 3.5, 1.2, [{"text": text, "size": 14, "color": MUTED}])


def roles(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 9, PAGES)
    title_block(slide, "Who can open what", "The menu hides the rest. A direct link outside the role is refused.")
    headers = ["Page", "Owner", "Catalog manager", "Order and support"]
    rows = [
        ["Dashboard", "Yes", "Yes", "Yes"],
        ["Products", "Yes", "Yes", ""],
        ["Categories", "Yes", "", ""],
        ["Orders", "Yes", "", "Yes"],
        ["Returns", "Yes", "", "Yes"],
        ["Customers", "Yes", "", "Yes"],
        ["Coupons", "Yes", "Yes", ""],
        ["Reports", "Yes", "", "Yes"],
        ["Settings", "Yes", "", ""],
    ]
    table_shape = slide.shapes.add_table(1 + len(rows), 4, Inches(0.62), Inches(1.85), Inches(12.1), Inches(4.95))
    table = table_shape.table
    table.columns[0].width = Inches(3.4)
    for col in range(1, 4):
        table.columns[col].width = Inches(2.9)
    for col, label in enumerate(headers):
        style_cell(table.cell(0, col), label, OLIVE_DEEP, CREAM, True, PP_ALIGN.LEFT if col == 0 else PP_ALIGN.CENTER, 13)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            fill = PAPER if r % 2 else CREAM
            color = OLIVE_DEEP if value == "Yes" else MUTED
            style_cell(table.cell(r, c), value, fill, INK if c == 0 else color, c == 0, PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)


def style_cell(cell, text, fill, color, bold, align, size=14):
    cell.text = text or "—"
    if not text:
        cell.text_frame.paragraphs[0].runs[0].text = "·"
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.14)
    cell.margin_right = Inches(0.08)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold or text == "Yes"
    run.font.color.rgb = color
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for old in tcPr.findall(qn(f"a:{edge}")):
            tcPr.remove(old)
        line = etree_line(edge)
        tcPr.append(line)


def etree_line(edge):
    from lxml import etree

    node = etree.Element(qn(f"a:{edge}"))
    node.set("w", "6350")
    fill = etree.SubElement(node, qn("a:solidFill"))
    color = etree.SubElement(fill, qn("a:srgbClr"))
    color.set("val", "E4DDD2")
    return node


def dashboard(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 10, PAGES)
    title_block(slide, "The morning view", "Every role lands here. The links onward appear only when that role can open them.")
    stats = [
        ("Revenue today", "Set against yesterday."),
        ("Orders today", "Set against yesterday."),
        ("Pending shipments", "Confirmed and packed, including any from before today."),
        ("Low stock", "Active pieces with one to five units left in a colour and size."),
    ]
    for index, (name, text) in enumerate(stats):
        x = 0.62 + index * 3.15
        rect(slide, x, 1.9, 3.0, 2.15, PAPER, LINE)
        rect(slide, x, 1.9, 3.0, 0.08, TERRA if index == 3 else OLIVE)
        add_text(slide, x + 0.18, 2.15, 2.65, 0.7, [{"text": name, "font": "Georgia", "size": 18}])
        add_text(slide, x + 0.18, 2.95, 2.65, 0.85, [{"text": text, "size": 13, "color": MUTED}])
    bullet_block(
        slide,
        0.62,
        4.35,
        12,
        2.4,
        [
            "A bar of sales for the last seven days.",
            "The latest orders, and an export of the order list.",
            "Low stock names the piece, the colour, and the size.",
        ],
        16,
    )


def catalog(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 11, PAGES)
    title_block(slide, "The catalog")
    rect(slide, 0.62, 1.85, 7.55, 4.85, PAPER, LINE)
    add_text(slide, 0.88, 2.05, 7, 0.4, [{"text": "Products", "font": "Georgia", "size": 22}])
    bullet_block(
        slide,
        0.88,
        2.6,
        7.0,
        3.8,
        [
            "Search, then filter by status, category, audience, and stock.",
            "Add a piece, or import a spreadsheet. A new piece starts in XS to XL.",
            "Edit name, description, fabric, care, price, and a dated sale price.",
            "Active, draft, or archived. Draft and archived stay off the shop.",
            "Up to eight colours and eight photographs. The first photo is the card.",
            "Stock by colour and size. Restocking emails anyone who asked.",
        ],
        15,
    )
    rect(slide, 8.4, 1.85, 4.3, 4.85, OLIVE_DEEP)
    add_text(slide, 8.65, 2.05, 3.9, 0.4, [{"text": "Categories", "font": "Georgia", "size": 22, "color": CREAM}])
    add_text(slide, 8.65, 2.55, 3.8, 0.4, [{"text": "Owner only.", "size": 13, "color": SAND}])
    bullet_block(
        slide,
        8.65,
        3.15,
        3.8,
        3.2,
        [
            "Name, a short line, and a portrait.",
            "Upper-body or lower-body size chart.",
            "How many pieces sit in each.",
            "New in stays a filter, not a category.",
            "The header, home, and shop read this list.",
        ],
        14,
        CREAM,
    )


def orders(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 12, PAGES)
    title_block(slide, "An order, from the list to the label")
    left = [
        "Search by number, name, email, or phone.",
        "Filter: all, confirmed, packed, shipped, delivered, cancelled.",
        "Open one order: items, SKU, GST, and the total.",
        "Move the status. The customer is emailed.",
        "A note that only staff can see.",
    ]
    right = [
        "Customer, address, and whether payment is paid or still due.",
        "Create a shipment with a courier and a tracking number.",
        "Print a shipping label, and print the GST invoice.",
        "Cancelling a packed order returns the stock and the coupon use.",
        "A paid Razorpay order is sent for refund.",
    ]
    bullet_block(slide, 0.62, 1.85, 6.0, 4.7, left, 16)
    bullet_block(slide, 6.9, 1.85, 5.8, 4.7, right, 16)


def desk(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 13, PAGES)
    title_block(slide, "Returns, customers, and coupons")
    columns = [
        (
            "Returns",
            [
                "Waiting, rupees still to refund, the return rate, and days to decide.",
                "Approve or reject.",
                "A refund returns stock and the Razorpay payment.",
                "An exchange is emailed for staff to arrange.",
                "A rejection puts the order back to delivered.",
            ],
        ),
        (
            "Customers",
            [
                "Search name, email, or phone.",
                "Everyone, or repeat buyers only.",
                "Orders, amount spent, and the last order.",
                "Guest checkouts appear by the email used.",
                "Block a person. They cannot sign in or order.",
            ],
        ),
        (
            "Coupons",
            [
                "Percent off, a fixed amount, or free shipping.",
                "A minimum order, a use limit, and dates.",
                "Active, scheduled, or expired.",
                "Open a code to edit it.",
                "The shop lists live codes in the bag and at checkout.",
            ],
        ),
    ]
    for index, (name, items) in enumerate(columns):
        x = 0.62 + index * 4.2
        rect(slide, x, 1.8, 4.0, 4.95, PAPER, LINE)
        rect(slide, x, 1.8, 4.0, 0.08, TERRA if index == 0 else OLIVE)
        add_text(slide, x + 0.22, 2.05, 3.55, 0.45, [{"text": name, "font": "Georgia", "size": 22}])
        bullet_block(slide, x + 0.22, 2.65, 3.55, 3.8, items, 13)


def reports(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 14, PAGES)
    title_block(slide, "Sales, tax, and how people paid", "All time, or the last 30 days. Cancelled orders stay out of the totals.")
    blocks = [
        ("Sales", "Net sales, order count, average order, GST collected, and refunds. Bars for category and for payment method."),
        ("GST", "Taxable value, CGST, SGST, and IGST, month by month. Karnataka splits CGST and SGST. Other states use IGST."),
        ("Payments", "UPI, cards, cash on delivery, and net banking, each as a share of sales."),
        ("The rate", "GST is already in the shelf price. Five percent up to ₹1,000 a piece. Twelve percent above that."),
    ]
    for index, (name, text) in enumerate(blocks):
        y = 1.8 + index * 1.25
        rect(slide, 0.62, y, 12.1, 1.12, PAPER, LINE)
        add_text(slide, 0.88, y + 0.14, 2.2, 0.8, [{"text": name, "font": "Georgia", "size": 18}], anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 3.2, y + 0.16, 9.2, 0.8, [{"text": text, "size": 15, "color": INK}], anchor=MSO_ANCHOR.MIDDLE)


def settings(prs):
    slide = blank(prs)
    chrome(slide, "The studio", 15, PAGES)
    title_block(slide, "What the owner can change", "Nothing reaches the shop until Save changes. The owner also invites staff.")
    groups = [
        ("The shop front", "Announcement, homepage words, hero photographs, the season picture, and the Women, Men, and Unisex tiles."),
        ("Money", "A store-wide sale, shipping rates, the free-shipping line, cash on delivery, and its fee and limit."),
        ("After the sale", "Return window in days, and the refund wording. Exchanges can be recorded here."),
        ("The paper", "Business name, GSTIN, prices inclusive of GST, and the studio address on the invoice and the footer."),
    ]
    for index, (name, text) in enumerate(groups):
        col = index % 2
        row = index // 2
        x = 0.62 + col * 6.35
        y = 1.9 + row * 2.4
        rect(slide, x, y, 6.1, 2.2, PAPER, LINE)
        add_text(slide, x + 0.28, y + 0.22, 5.55, 0.4, [{"text": name, "font": "Georgia", "size": 20}])
        add_text(slide, x + 0.28, y + 0.8, 5.55, 1.15, [{"text": text, "size": 15, "color": MUTED}])


def letters(prs):
    slide = blank(prs)
    chrome(slide, "Both sides", 16, PAGES)
    title_block(slide, "The letters the site sends", "These go out when mail is configured.")
    notes = [
        ("Account", "A password reset. The link expires in 30 minutes."),
        ("Back in stock", "A note that the request was kept, then a second letter when that colour and size returns, with a link to the piece."),
        ("The order", "Confirmed, packed, shipped, out for delivery, delivered, and cancelled."),
        ("The return", "Requested, approved, exchange approved, declined, and returned."),
    ]
    for index, (name, text) in enumerate(notes):
        y = 1.8 + index * 1.25
        rect(slide, 0.62, y, 0.12, 1.08, TERRA if index % 2 else OLIVE)
        add_text(slide, 1.05, y + 0.05, 3, 0.4, [{"text": name, "font": "Georgia", "size": 20}])
        add_text(slide, 4.15, y + 0.08, 8.4, 0.9, [{"text": text, "size": 16}])


def close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NIGHT
    rect(slide, 0, 0, 0.18, 7.5, OLIVE)
    add_text(slide, 0.85, 1.7, 11, 0.35, [{"text": "THE WALK", "size": 13, "color": SAND, "spacing": 2.4}])
    add_text(
        slide,
        0.85,
        2.15,
        11.5,
        1.3,
        [
            {"text": "One customer.", "font": "Georgia", "size": 40, "color": CREAM, "after": 0},
            {"text": "Then one owner.", "font": "Georgia", "size": 40, "color": CREAM},
        ],
    )
    add_text(
        slide,
        0.85,
        4.3,
        10.5,
        1.4,
        [
            {"text": "Show the shop from the homepage to a placed order.", "size": 18, "color": SAND, "after": 8},
            {"text": "Sign out. Open the studio, and show that same order moving.", "size": 18, "color": SAND},
        ],
    )
    add_text(slide, 0.85, 6.55, 8, 0.3, [{"text": "Earthen", "font": "Georgia", "size": 16, "color": OLIVE, "italic": True}])


def notes(prs):
    slide = blank(prs)
    rect(slide, 0, 0, 13.333, 0.08, TERRA)
    add_text(slide, 0.62, 0.32, 8, 0.3, [{"text": "FOR US  ·  SKIP THIS SLIDE IN THE ROOM", "size": 12, "color": TERRA, "bold": True, "spacing": 1.4}])
    add_text(slide, 0.62, 0.68, 12, 0.6, [{"text": "Say these if they come up", "font": "Georgia", "size": 32}])
    items = [
        ("Sign-in", "Email and password. A one-time code is not part of this build."),
        ("Messages", "The customer can choose SMS or WhatsApp. The letters that go out today are email."),
        ("Contact and the list", "Both confirm on the page. They do not yet send to the studio or store the address."),
        ("Payment", "UPI, cards, and net banking need Razorpay. Without it, the page says no payment is taken. Cash on delivery still places an order."),
        ("Two settings", "Allow exchanges is saved, and the return form still offers both. Homepage banners are a staff list. The live homepage uses the words and the pictures."),
    ]
    for index, (name, text) in enumerate(items):
        y = 1.55 + index * 1.02
        add_text(slide, 0.62, y, 2.4, 0.85, [{"text": name, "font": "Georgia", "size": 16}], anchor=MSO_ANCHOR.TOP)
        add_text(slide, 3.15, y, 9.4, 0.9, [{"text": text, "size": 15, "color": INK}])


if __name__ == "__main__":
    print(build())
